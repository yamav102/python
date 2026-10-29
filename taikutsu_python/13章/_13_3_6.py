#! python3
# P356
# _13_3_6.py
# Word文書を書き込む

import docx
# Word新規作成
doc = docx.Document()
print(doc.add_paragraph('Hello world!'))
# <docx.text.paragraph.Paragraph object at 0x000001FFCC45B230>
para_obj1 = doc.add_paragraph('これは第2段落です。')
para_obj2 = doc.add_paragraph('これは第3段落です。')
print(para_obj1.add_run(' これは第2段落に追加したテキストです。'))
# <docx.text.run.Run object at 0x00000214AC25C8A0>
para_obj1.insert_paragraph_before('これは第2段落の手前に挿入された段落です。')

# 段落のインデックスを指定してinsert する（スタイル指定付）
doc.paragraphs[0].insert_paragraph_before(
    'これは後から挿入されたタイトルです', 
    'Title'
    )
doc.save('hellloworld.docx')
