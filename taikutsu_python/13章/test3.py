#! python3
# P356
# Word文書に定義しえたるスタイル名を調べる
# test3.py

from docx.document import Document
import docx
def print_style(doc: Document)->None:
    # doc = docx.Document('demo.docx')
    for s in doc.styles:
        print(s.name)

    print('------------------------------------')    

    for p in doc.paragraphs:
        for r in p.runs:
            # print((r.text or 'None') + '\t' + (r.style.name or 'None'))
            txt = r.text or 'None'
            if r.style and r.style.name:
                style_name = r.style.name
            else:
                style_name = 'None'            
            print(txt.ljust(30,'-') + ': ' + style_name)
def main()->None:            
    doc = docx.Document('demo.docx')
    print_style(doc)
    doc = docx.Document('restyled.docx')
    print_style(doc)

    print('done')

if __name__ == '__main__':
    main()