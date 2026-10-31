#! python3
# P358
# _13_3_7.py
# 見出しを追加する
import docx

doc = docx.Document()
doc.add_heading('表題', 0) # 見出し文字列、Level 0～9
doc.add_heading('Header 1', 1)
doc.add_heading('Header 2', 2)
doc.add_heading('Header 3', 3)
doc.add_heading('Header 4', 4)
doc.add_heading('Header 5', 5)
doc.add_heading('Header 6', 6)
doc.add_heading('Header 7', 7)
doc.add_heading('Header 8', 8)
doc.add_heading('Header 9', 9)

print(doc.paragraphs[0].text)
print(doc.paragraphs[1].text)
print(doc.paragraphs[2].text)
print(doc.paragraphs[3].text)
print(doc.paragraphs[4].text)
print(doc.paragraphs[5].text)
print(doc.paragraphs[6].text)
print(doc.paragraphs[7].text)
print(doc.paragraphs[8].text)
print(doc.paragraphs[9].text)

# 表題
# Header 1
# Header 2
# Header 3
# Header 4
# Header 5
# Header 6
# Header 7
# Header 8
# Header 9

doc.save('headings.docx')

