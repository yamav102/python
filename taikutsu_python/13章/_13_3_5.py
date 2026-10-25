#! python3
# P355
# _13_3_5.py
# Runの属性

import docx
from pathlib import Path
wkdir = Path(__file__).parent
doc = docx.Document(str(wkdir / 'demo.docx'))
print(doc.paragraphs[0].text) # Document Title
print(doc.paragraphs[0].style) # _ParagraphStyle('Title') id: 2733980373344 '表題'
doc.paragraphs[0].style = 'Normal'
print(doc.paragraphs[0].style) # _ParagraphStyle('Normal') id: 2733973253792 '標準'

print(doc.paragraphs[1].text) # A plain paragraph with some bold and some italic
para1 = doc.paragraphs[1]
print(
    (para1.runs[0].text,
     para1.runs[1].text,
     para1.runs[2].text,
     para1.runs[3].text,)
) # ('A plain paragraph with some ', 'bold', ' and some ', 'italic')

para1.runs[0].style = 'Quote Char' # 引用文
# 'bold', 'italic' に underline 
para1.runs[1].underline = True
para1.runs[3].underline = True

doc.save(str(wkdir /'restyled.docx'))





