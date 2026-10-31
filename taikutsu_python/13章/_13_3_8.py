#! python3
# P359
# _13_3_8.py
# 改行と改ページを追加する
import docx
doc = docx.Document()
doc.add_paragraph('これは1ページ目')
# 改ページ
doc.add_page_break()
doc.add_paragraph('これは2ページ目')
doc.save('twoPage.docx')