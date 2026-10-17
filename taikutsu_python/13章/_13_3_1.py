#! python3
# P349
# 13.3 Word文書
# _13_3_1 Word文書を読み込む
import sys
# print(sys.path)
# for p in sys.path:
    # print(p)
import docx
def main()->None:
    doc = docx.Document('demo.docx')
    print(len(doc.paragraphs))
    print(doc.paragraphs[0].text)
    print(doc.paragraphs[1].text)
    print(len(doc.paragraphs[1].runs))
    print(doc.paragraphs[1].runs[0].text)
    print(doc.paragraphs[1].runs[1].text)
    print(doc.paragraphs[1].runs[2].text)
    print(doc.paragraphs[1].runs[3].text)
    for rns in doc.paragraphs[1].runs:
        print(rns.text)
# 7        
# Document Title
# A plain paragraph with some bold and some italic
# 4
# A plain paragraph with some 
# bold
#  and some 
# italic
# A plain paragraph with some 
# bold
#  and some 
# italic        

if __name__ == '__main__':
    main()
