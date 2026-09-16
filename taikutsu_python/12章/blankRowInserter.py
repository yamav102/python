#! python3
# 12.14.2 空行を挿入する
# P332
# blankRowInserter.py
# コマンドライン引数から 2つの整数とファイル名を受け取る
# arg1:N, arg2:M , fileName
# N行に M行分の空行を挿入する
# excel がインストールされていない環境想定
# 新しいシートにN行書き込んだ後に、M行足した行に続きを書き込む
import argparse
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from typing import cast
def main(N:int, M:int, filepath:str)->None:
    if N < 1 or M < 1:
        raise ValueError(f'挿入指定行 > 0, 挿入行数 > 0')
    wb = openpyxl.load_workbook(filename=filepath) # , data_only=False)既定
    # 最も左にあるシートを対象とする
    ws = cast(Worksheet, wb.worksheets[0]) 
    mxrw = ws.max_row
    if mxrw < N:
        raise ValueError(
            f'挿入指定行 {N} は、最大使用行 {mxrw} より大きく、'
            + '行挿入されませんでした。')
    ws_result = wb.create_sheet(title='hoge', index=0)    
    for rw in range(ws.min_row, ws.max_row+1):
        for cl in range(ws.min_column, ws.max_column+1):
            if rw < N:                
                ws_result.cell(rw, cl).value = ws.cell(rw, cl).value                
            else:
                ws_result.cell(rw+M, cl).value = ws.cell(rw, cl).value
        
    wb.save(filepath)    
    wb.close()            
    print('完了：' + filepath)

if __name__ == '__main__':
    parser = argparse.ArgumentParser(
        description='N行目にM行の空行を挿入します。',
        epilog='py blankRowInserter.py N M filename.xlsx'
    )
    parser.add_argument(
        'targetRow',
        type=int,
        help='空行挿入する行番号'
    )
    parser.add_argument(
        'insertRows',
        type=int,
        help='挿入する行数'
    )
    parser.add_argument(
        'filename',
        type=str,
        help='対象ファイル名'
    )    
    args = parser.parse_args()
    main(args.targetRow, args.insertRows, args.filename)
