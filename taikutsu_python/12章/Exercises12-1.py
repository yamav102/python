#! python3
# P330 Exercises12-1.py
# 12.13　演習問題
# 12-1
## load_workbook() は wbオブジェクト を返す
# 12-2
## wb.get_sheet_name(’hoge’) は 名前を指定してシートオブジェクトを返す古い書き方
## workbook['hoge'] が今の書き方
## wb.sheetnames は シート名をlist で返す
# 12-3 
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
wb = openpyxl.Workbook()
ws1:Worksheet = wb.create_sheet(title='Sheet1',index=1)
ws2 = wb.create_sheet(title='Sheet2',index=2)
ws3 = wb.create_sheet(title='Sheet3',index=3)
wshoge = wb.create_sheet(title='hoge',index=0)
ws16 = wb.create_sheet(title='Sheet16',index=16) # エラーにならない。最右に追加される
# [<Worksheet "hoge">, <Worksheet "Sheet">, 
# <Worksheet "Sheet1">, <Worksheet "Sheet2">, <Worksheet "Sheet3">, <Worksheet "Sheet16">]
print(len(wb.worksheets))
# 6
print(wb['Sheet1'])
# <Worksheet "Sheet1">
print(wb.index(ws1)) # 2 index は wb が管理している属性

# 12-4
wb.active = wb['Sheet1']
ws = wb.active
if ws is not None:
    print('activesheetは、' + ws.title) # activesheetは、Sheet1
print(wb['hoge'].title) # hoge
print(wb.worksheets[1].title) # Sheet

# 12-5
# wb['sheet1']['c5'].Value = 'Hello' シート名の大文字小文字は区別される→ERR
# wb['Sheet1']['c5'].Value = 'Hello' Value属性はない .value である。
wb['Sheet1']['c5'] = 'Hello'
ws = wb['Sheet1']
print(ws['c5'].value)
print(wb['Sheet1']['c5'].value)
print(ws['c5'])
print(wb['Sheet1']['c5'])
# Hello
# Hello
# <Cell 'Sheet1'.C5>
# <Cell 'Sheet1'.C5>
# from openpyxl.workbook.workbook import Workbook
# help(Workbook)
# help(Worksheet)
# 12-6
wb['Sheet1']['c5'].value = 'Hello World'
print(ws['c5'].value) # Hello World

# 12-7
rng = ws['d6']
print(type(rng)) # <class 'openpyxl.cell.cell.Cell'>
print(rng.row, rng.column) # 6 4

# 12-8
print(ws.max_row, ws.max_column, ws.min_row, ws.min_column)
# 6 4 5 3
print(ws.cell(5, 3).coordinate) # C5
print(type(ws.max_row), type(ws.max_column))
# <class 'int'> <class 'int'>
# from openpyxl.cell.cell import Cell

# 12-9
print(ws['M1'].column) # 13
from openpyxl.utils import column_index_from_string
print(column_index_from_string('M')) # 13

# 12-10
from openpyxl.utils import get_column_letter
print(get_column_letter(14)) # N

# 12-11
tpl = ws['a1:f1']
# print(tpl, type(tpl))
# (
# (
# <Cell 'Sheet1'.A1>, 
# <Cell 'Sheet1'.B1>, 
# <Cell 'Sheet1'.C1>, 
# <Cell 'Sheet1'.D1>, 
# <Cell 'Sheet1'.E1>, 
# <Cell 'Sheet1'.F1>
# ),
# ) ← 複数行を想定した形式で出力されている。一行だけ欲しい場合は tple[0]
# <class 'tuple'>
print(tpl[0])
# (<Cell 'Sheet1'.A1>, <Cell 'Sheet1'.B1>, <Cell 'Sheet1'.C1>, <Cell 'Sheet1'.D1>, <Cell 'Sheet1'.E1>, <Cell 'Sheet1'.F1>)
# tpl = ws['a1':'f1']   # この記法は、pythonの スライス風にも書けるよ、というために用意されている（だけ）
# print(tpl, type(tpl))

# 12-12, 12-13
ws['F1'] = 100
ws['F2'] = 200
ws['F3'] = '=SUM(F1:F2)'
ws['F4'] = '=200+800'
print(ws['F1'].value)
print(ws['F2'].value)
print(ws['F3'].value)
print(ws['F4'].value)
# 100
# 200
# =SUM(F1:F2) openpyxl は計算はしない。
# =200+800
from pathlib import Path
xlpath = Path(__file__).parent / 'exersises12.xlsx'
wb.save(xlpath)
def recalc_save(xlpath:str)->None:
    from win32com.client import DispatchEx # 新規excel
    xlapp = DispatchEx('excel.application')
    # COM が欲しがる文字列の絶対パス
    xlwb = xlapp.Workbooks.Open(str(Path(xlpath).resolve())) 
    xlwb.Save() # ここで計算結果がキャッシュされる
    xlwb.Close(SaveChanges:=True)
    xlapp.Quit()

recalc_save(str(xlpath))
# 計算結果用（数式は消える。上書き保存しないこと）
wb_val = openpyxl.load_workbook(xlpath, data_only=True)
# wb_val.save(xlpath) saveすると数式消えます。
print(wb_val['Sheet1']['F3'].value)
# 300
print(wb_val['Sheet1']['F4'].value)
# 1000

# COMで開いたエクセルから直接計算結果を得る
def get_value(xlpath:str)->None:
    from win32com.client import Dispatch
    xlapp = Dispatch('excel.application')
    # COM が欲しがる文字列の絶対パス
    xlwb = xlapp.Workbooks.Open(str(Path(xlpath).resolve())) 
    ws = xlwb.Worksheets('Sheet1')
    print('COM: ' + str(ws.Range('f3').Value))
    print('COM: ' + str(ws.Range('f4').Value))
    # COM: 300.0
    # COM: 1000.0    
    xlwb.Close(SaveChanges:=False)
    xlapp.DisplayAlerts = False
    xlapp.Quit()

get_value(str(xlpath))

# 12-14
# セルの計算結果は、一度エクセルに計算させて保存したファイルを data_only
# で開く必要がある。単に openpyxl で保存して、data_only で load するだけでは
# だめ。⇒ test9.py で検証した。

# 12-15 行の高さを 100 にする。
from typing import cast
wb = openpyxl.Workbook()
ws = cast(Worksheet, wb.active)
ws.row_dimensions[5].height = 100 # point
ws.column_dimensions['A'].width = 15 # 半角の文字数？らしいが、実際に試してみると入りきらなかった。
# ↑ COM から columnwidth で設定すれば、VBAでの設定と同じ幅になる。
# openpyxl と vba の列幅設定の結果は同じにならない。
# COM からなら autofit も使える。
xlpath = Path(__file__).parent / 'xlRowHeight.xlsx'
xlpath = str(xlpath.resolve())
wb.save(xlpath)
from win32com.client import Dispatch
from typing import Any, Tuple # 何でもあり、という事
def xlopen(xlpath:str)-> Tuple[Any, Any]: # CDispatch を指定しても、補完が利くわけでもないので、Any を当てて置く。
    xlapp = Dispatch('excel.application')
    xlapp.Visible = True
    return (xlapp, xlapp.Workbooks.Open(xlpath))
xlapp, xlwb = xlopen(xlpath)    
print(type(xlwb)) # <class 'win32com.client.CDispatch'>
# point ピクセル変換の過程で端数が切り捨てられたものの積をとったりするので、
# 誤差を含む値が戻るのは仕方がない。
print(xlwb.Worksheets(1).Rows(5).Height) # 99.75
xlwb.Close()
xlapp.Quit()

# # 12-16
wb = openpyxl.Workbook()
xlpath = Path(__file__).parent / 'hiddencolumn.xlsx'
xlpath = str(xlpath.resolve())
ws = cast(Worksheet, wb.active)
ws.column_dimensions['C'].hidden = True
wb.save(xlpath)
xlapp, xlwb = xlopen(xlpath)
xlwb.Close()
xlapp.Quit()

# 12-17
# openpyxl 2.4.1 は 画像、グラフ、ウィンドウ枠の固定を読み込まない
print(openpyxl.__version__) # 3.1.5

# 12-18
# ウィンドウ枠の固定
wb = openpyxl.Workbook()
ws = cast(Worksheet, wb.active)
ws.sheet_view.topLeftCell='a1'
ws.freeze_panes='a2'
xlpath = Path(__file__).parent / 'freezepanes.xlsx'
xlpath = str(xlpath.resolve())
wb.save(xlpath)
wb.close()
xlapp, xlwb = xlopen(xlpath)
xlwb.Close()
xlapp.Quit()

# 12-19
# グラフ
# 1) refオブジェクトを作って、
# 2) チャーとオブジェクトを作って
# 3) 2) に 1) を add_data する
from openpyxl.chart import Reference
from openpyxl.chart import BarChart
from openpyxl.cell.cell import Cell
# refernce を作って、BarChartへ add_dataする
wb = openpyxl.Workbook()
ws = cast(Worksheet, wb.active)
for i in range(1, 5):
    cast(Cell,ws.cell(i, 1)).value = i
ref = Reference(
    ws, min_row=1, min_col=1, max_row=4, max_col=1
)
grf = BarChart()
grf.add_data(ref)
# 軸目盛を表示
grf.x_axis.delete = False 
grf.y_axis.delete = False
grf.title = 'HOGE'
grf.height = 10.5 # cm
grf.width = 10 # cm
# 1cm ≒ 28.35pt
chart_h_pt = grf.height * 28.35 # グラフエリアのサイズ point
plot_top_pt = 20 # プロットエリアの top
plot_h_pt = 250 # プロットエリアの height

# グラフ plotarea の位置とサイズ指定
from openpyxl.chart.layout import Layout, ManualLayout
# grf.height = chart_h_pt / 28.35 # 290/28.35 →だいたい 10cm
# プロットエリアのサイズ
grf.layout = Layout(
    manualLayout=ManualLayout(
        layoutTarget="inner",
        xMode="edge",
        yMode="edge",
        x=0.10,
        # プロットエリアの top
        # /チャートエリアの高さの比で渡す
        y=plot_top_pt / chart_h_pt,
        # 幅の割合
        w=0.80,
        # Height 
        # プロットエリアの高さ/チャートエリアの高さの比で渡す
        h=plot_h_pt / chart_h_pt, 
    )
)

ws.add_chart(grf, 'C2')
xlpath = Path(__file__).parent / 'graph.xlsx'
xlpath = str(xlpath.resolve())
wb.save(xlpath)
wb.close()

xlapp, xlwb = xlopen(xlpath)

