#! python3
# 12.14.2 空行を挿入する
# P332
# blankRowInserter.py
# コマンドライン引数から 2つの整数とファイル名を受け取る
# arg1:N, arg2:M , fileName
# N行に M行分の空行を挿入する
# excel がインストールされていない環境想定
# Worksheet.insert_rows(N, M) を使う。
import argparse
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from typing import cast
import sys
def main(N:int, M:int, filepath:str)->None:
    if N < 1 or M < 1:
        raise ValueError(f'挿入指定行 > 0, 挿入行数 > 0')
    wb = openpyxl.load_workbook(filename=filepath)
    # 最も左にあるシートを対象とする
    ws = cast(Worksheet, wb.worksheets[0])
    mxrw = ws.max_row
    mxclm = ws.max_column
    # シートにデータが無い場合の情報出力
    if all([mxrw==1, mxclm==1]):
        if ws['A1'].value in (None, ''):
            print(
                'データが無いシートなので、挿入は行いませんでした。'
                ,filepath
                )            
            sys.exit(0) # パイプラインで失敗にしたいなら 1
    # 挿入指定行がデータ行より大きい場合エラー
    if mxrw < N:
        raise ValueError(
            f'挿入指定行 {N} は、最大使用行 {mxrw} より大きく、'
            + '行挿入されませんでした。')
    
    # openpyxl の制限（この API を使う以上避けられない）
    # 数式・名前付き範囲・グラフ・テーブルは行挿入に追従しないことが多い。
    # 結合セルがあるとずれやすい（結合範囲を先に shift する必要がある）。
    # 挿入行の高さ・スタイルは元の行から自動では十分にコピーされない。    
    ws.insert_rows(idx=N, amount=M)
    wb.save(filepath)    
    wb.close()            
    print('完了：' + filepath)

if __name__ == '__main__':
    parser = argparse.ArgumentParser(
        description='N行目にM行の空行を挿入します。',
        epilog='py blankRowInserter.py N M filename.xlsx'
    )
    parser.add_argument(
        'targetRow',
        type=int,
        help='空行挿入する行番号'
    )
    parser.add_argument(
        'insertRows',
        type=int,
        help='挿入する行数'
    )
    parser.add_argument(
        'filename',
        type=str,
        help='対象ファイル名'
    )    
    args = parser.parse_args()
    main(args.targetRow, args.insertRows, args.filename)
