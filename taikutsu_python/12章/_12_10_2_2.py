#! python3
# _12_10_2_2.py
# P325
# セルの結合と解除
import openpyxl
from pathlib import Path
from openpyxl.worksheet.worksheet import Worksheet
import win32com.client
def openxl(xlpath: str)->None:
    xlApp = win32com.client.Dispatch('Excel.Application')
    xlApp.Visible = True
    xlApp.Workbooks.open(xlpath)
def main():
    xlpath = Path(__file__).parent / 'merge_cell.xlsx'
    wb = openpyxl.load_workbook(xlpath)
    sheet = wb.active
    assert isinstance(sheet, Worksheet)
    sheet.unmerge_cells('a1:d3') # 結合解除
    sheet.unmerge_cells('c5:d5')
    wb.save(xlpath)
    openxl(xlpath=str(xlpath))
if __name__ == '__main__':
    # print(__name__)
    main()