#! python3
# 12.14.1
# multiplicationTable.py
# P331
# 掛け算の表を作成する
# コマンドラインから N を受け取り、N × N の表をExcelシートに作成する。
import argparse
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from win32com.client import DispatchEx
from typing import Tuple, Any, cast
from pathlib import Path

def openxl(fpath:str)->Tuple[Any, Any]:
    xlapp = DispatchEx('excel.application')
    wb = xlapp.Workbooks.Open(fpath)
    return (xlapp, wb)

def main(n: int, xlpath: str)->None:
    # workbook生成
    wb = openpyxl.Workbook()
    ws = cast(Worksheet, wb.active)
    # 見出し
    font_headerstyle = Font(bold=True, color='FFFF0000') # 先頭のFFは不透明度
    for i in range(1, n+1):        
        topheader = cast(Cell, ws.cell(1, i+1))
        leftheader = cast(Cell, ws.cell(i+1, 1))
        topheader.value = i
        leftheader.value = i
        topheader.font = font_headerstyle
        leftheader.font = font_headerstyle

        # データ
        for cl in range(2, n+2):
            cast(Cell, ws.cell(i+1, cl)).value = (
                # str(i) + ': ' + str(cl-1)
                i * (cl-1)
            )            
    wb.save(xlpath)
    wb.close

    # wbopen
    xlapp, xlwb = openxl(fpath=xlpath)
    xlapp.Visible = True

if __name__ == '__main__':
    parser = argparse.ArgumentParser(
        description='N * N の表を作成します。',
        epilog="py multiplicationTable.py N"
    )
    parser.add_argument(
        'number',
        type=int,
        help='整数'
        )
    args = parser.parse_args()
    xlpath = Path(__file__).parent / 'multiplicationTable.xlsx'
    xlpath = str(xlpath.resolve())
    main(args.number, xlpath=xlpath)