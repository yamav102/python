#! python3
# P332
# 12.14.3 行と列の入れ替え
# エクセルファイルパスを渡し、
# 二重の for ループで データを読み込み、行列を入れ替えて出力する
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import Cell, MergedCell
# from typing import cast
import argparse
from pathlib import Path
# 上書きしないファイルパスを返す
def mk_save_name(xlpath:str)->str:
    p = Path(xlpath)
    cand = p # cand:candidate>候補    
    i = 1
    while cand.exists():        
        # with_name は parent はそのままで、
        # ファイル名部分のみ差し替えるメソッド
        cand = p.with_name(f'{p.stem}({i}){p.suffix}')
        i += 1
    return str(cand)

def main(xlpath:str)->None:
    wb = openpyxl.load_workbook(xlpath, data_only=True)
    ws:Worksheet = wb.worksheets[0] # 最左シート
    wsn:Worksheet = wb.create_sheet(index=0)
    # ws使用範囲をループしながらwsn に値を転記
    for rw in range(ws.min_row, ws.max_row+1):
        for clm in range(ws.min_column, ws.max_column+1):
            # 型チェッカー対策。実行だけなら次の1行で足りる
            # wsn.cell(row=clm, column=rw).value = ws.cell(row=rw, column=clm).value            
            src = ws.cell(row=rw, column=clm)
            if isinstance(src, MergedCell):
                val = None
            else:
                val = src.value
            dst = wsn.cell(row=clm, column=rw)
            assert isinstance(dst, Cell)
            dst.value = val
        
    shname = ws.title
    wb.remove(ws) # シート削除
    wsn.title = shname
    wb.save(mk_save_name(xlpath))
    wb.close()
    print('done')
if __name__ == '__main__':
    parser = argparse.ArgumentParser(
        description='行列を入れ替えます',
        epilog='py _12_14_3.py hoge.xlsx'
    )
    parser.add_argument(
        # --xlpath とすると名前付引数になる。
        # --または- を付けた場合は名前付で指定しないとエラーになる。
        'xlpath', 
        type=str,
        help='行列入れ替えたいxlsxファイルパス'
    )
    args = parser.parse_args()
    main(args.xlpath)