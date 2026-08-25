#! python3
# P325 _12_10_2.py
# セルの結合と解除

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from pathlib import Path
import time
import win32com.client
def main():
    wb = openpyxl.Workbook()
    sheet = wb.active
    assert isinstance(sheet, Worksheet)
    sheet.merge_cells('a1:d3')
    sheet['a1'] = 'Twelve cells merged together.'
    sheet.merge_cells('c5:d5')
    sheet['c5'] = 'Two merged cells.'
    xlpath = Path(__file__).parent / 'merge_cell.xlsx'
    wb.save(xlpath)
    time.sleep(.1)
    xlApp = win32com.client.Dispatch('Excel.Application')
    xlApp.Visible = True
    xlApp.Workbooks.Open(xlpath)


if __name__ == '__main__':
    main()